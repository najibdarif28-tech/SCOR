from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


# Moody's-inspired palette
MOODYS_BLUE = RGBColor(0x00, 0x35, 0x7A)
MOODYS_LIGHT_BLUE = RGBColor(0xE9, 0xF0, 0xFA)
MOODYS_ACCENT = RGBColor(0x36, 0x7D, 0xB8)
DARK_TEXT = RGBColor(0x21, 0x2B, 0x36)
MUTED_TEXT = RGBColor(0x5A, 0x67, 0x75)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def add_header(slide, title_text):
    # Top title bar
    top_bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.7)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = MOODYS_BLUE
    top_bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.35), Inches(0.12), Inches(9.5), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = WHITE

    brand_box = slide.shapes.add_textbox(Inches(10.0), Inches(0.12), Inches(3.0), Inches(0.45))
    bp = brand_box.text_frame.paragraphs[0]
    bp.text = "MOODY'S ANALYTICS"
    bp.font.size = Pt(13)
    bp.font.bold = True
    bp.font.color.rgb = WHITE
    bp.alignment = PP_ALIGN.RIGHT


def add_footer(slide, footer_text="Confidential | Internal Use Only"):
    footer = slide.shapes.add_textbox(Inches(0.35), Inches(7.2), Inches(12.6), Inches(0.25))
    p = footer.text_frame.paragraphs[0]
    p.text = footer_text
    p.font.size = Pt(10)
    p.font.color.rgb = MUTED_TEXT


def add_bullets(slide, bullets, left=0.7, top=1.2, width=12.0, height=5.5):
    body = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = body.text_frame
    tf.clear()
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(20 if i == 0 else 17)
        p.font.color.rgb = DARK_TEXT
        if i == 0:
            p.font.bold = True
    for p in tf.paragraphs[1:]:
        p.level = 1
        p.font.size = Pt(16)


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "SCOR Managing Agency | Contractual & Relationship Situation")

    hero = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.8), Inches(2.1)
    )
    hero.fill.solid()
    hero.fill.fore_color.rgb = MOODYS_LIGHT_BLUE
    hero.line.color.rgb = MOODYS_ACCENT

    tbox = slide.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(11.2), Inches(1.4))
    p = tbox.text_frame.paragraphs[0]
    p.text = "Executive summary of contractual position, use case, and latest client topics"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = MOODYS_BLUE

    p2 = tbox.text_frame.add_paragraph()
    p2.text = "Prepared from agreement renewals (2023-2025), amendment, and latest client correspondence"
    p2.font.size = Pt(16)
    p2.font.color.rgb = MUTED_TEXT

    info = slide.shapes.add_textbox(Inches(0.95), Inches(4.1), Inches(7.8), Inches(1.2))
    ip = info.text_frame.paragraphs[0]
    ip.text = "Client: SCOR Managing Agency Limited (formerly The Channel Managing Agency Limited)"
    ip.font.size = Pt(16)
    ip.font.color.rgb = DARK_TEXT
    ip2 = info.text_frame.add_paragraph()
    ip2.text = "Primary contact: Jeremy Wang (Head of Capital Modelling)"
    ip2.font.size = Pt(16)
    ip2.font.color.rgb = DARK_TEXT

    # Icon-style callouts
    labels = [("C", "Contract"), ("U", "Use Case"), ("T", "Topics")]
    for i, (letter, label) in enumerate(labels):
        x = 8.9 + i * 1.35
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(4.2), Inches(0.8), Inches(0.8))
        circle.fill.solid()
        circle.fill.fore_color.rgb = MOODYS_ACCENT
        circle.line.color.rgb = MOODYS_BLUE
        cp = circle.text_frame.paragraphs[0]
        cp.text = letter
        cp.font.size = Pt(20)
        cp.font.bold = True
        cp.font.color.rgb = WHITE
        cp.alignment = PP_ALIGN.CENTER

        lbox = slide.shapes.add_textbox(Inches(x - 0.05), Inches(5.05), Inches(1.0), Inches(0.3))
        lp = lbox.text_frame.paragraphs[0]
        lp.text = label
        lp.font.size = Pt(11)
        lp.font.color.rgb = MUTED_TEXT
        lp.alignment = PP_ALIGN.CENTER

    add_footer(slide)


def add_executive_summary_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "1) Executive Summary")
    bullets = [
        "Current situation",
        "SCOR is renewing annually under the long-running Moody's agreement framework.",
        "The 2025 renewal fee is GBP 53,950 (+3.0% vs 2024), below standard increase range referenced by sales (6-10%).",
        "Contract entity/name and address have been formally updated to SCOR Managing Agency Limited, 8 Bishopsgate.",
        "Client preference is a one-year renewal (vs multi-year), while Moody's proposed limited uplift continuity for multi-year.",
        "Recent conversations focus on operational simplification and parameter governance in Scenario Generator usage.",
    ]
    add_bullets(slide, bullets)
    add_footer(slide)


def add_contract_table_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "2) Contractual Position (Summary Table)")

    rows, cols = 6, 5
    table = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(1.15), Inches(12.3), Inches(4.7)).table
    headers = ["Document", "Effective / Renewal Date", "Agreement Ref", "Key Point", "Financial Impact"]
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = MOODYS_BLUE
        para = cell.text_frame.paragraphs[0]
        para.font.bold = True
        para.font.color.rgb = WHITE
        para.font.size = Pt(12)

    rows_data = [
        [
            "Base Order Form",
            "23 Dec 2016",
            "00073841.0",
            "Framework remains governing baseline for current services",
            "Legacy baseline",
        ],
        [
            "Renewal Notification",
            "23 Dec 2023",
            "00073841.7",
            "Annual renewal under existing services",
            "GBP 50,364",
        ],
        [
            "Renewal Notification",
            "23 Dec 2024",
            "00073841.8",
            "Service continuity maintained",
            "GBP 52,379",
        ],
        [
            "Renewal Notification",
            "23 Dec 2025",
            "00073841.9",
            "Annual fee reminder and services retained",
            "GBP 53,950",
        ],
        [
            "Amendment 1 to Order Form",
            "23 Dec 2025",
            "00073841.9",
            "Client renamed/address updated + sanctions clause refresh",
            "No explicit fee change",
        ],
    ]

    for r, row in enumerate(rows_data, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xF8, 0xFA, 0xFD) if r % 2 == 1 else RGBColor(0xEE, 0xF3, 0xFA)
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(11)
            para.font.color.rgb = DARK_TEXT

    note = slide.shapes.add_textbox(Inches(0.65), Inches(6.05), Inches(12.0), Inches(0.9))
    np = note.text_frame.paragraphs[0]
    np.text = (
        "Contractual note: amendment confirms legal identity transition from The Channel Managing Agency Limited to "
        "SCOR Managing Agency Limited and updates the licensed site to 8 Bishopsgate."
    )
    np.font.size = Pt(12)
    np.font.color.rgb = MUTED_TEXT
    add_footer(slide)


def add_graph_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "3) Commercial Trend (Annual Renewal Fees)")

    chart_data = CategoryChartData()
    chart_data.categories = ["2023", "2024", "2025"]
    chart_data.add_series("Annual Fee (GBP)", (50364, 52379, 53950))

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS, Inches(0.8), Inches(1.3), Inches(7.2), Inches(4.8), chart_data
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.series[0].format.line.color.rgb = MOODYS_BLUE
    chart.value_axis.has_major_gridlines = True
    chart.category_axis.has_major_gridlines = False
    chart.chart_title.has_text_frame = True
    chart.chart_title.text_frame.text = "Renewal fees (GBP)"

    callout = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.35), Inches(1.55), Inches(4.1), Inches(3.8)
    )
    callout.fill.solid()
    callout.fill.fore_color.rgb = MOODYS_LIGHT_BLUE
    callout.line.color.rgb = MOODYS_ACCENT
    tf = callout.text_frame
    tf.clear()
    p0 = tf.paragraphs[0]
    p0.text = "Key observations"
    p0.font.size = Pt(16)
    p0.font.bold = True
    p0.font.color.rgb = MOODYS_BLUE

    for line in [
        "2024 vs 2023: +GBP 2,015 (+4.0%)",
        "2025 vs 2024: +GBP 1,571 (+3.0%)",
        "Sales note references standard market uplift of 6-10%",
        "Current renewal kept to lower increase level",
    ]:
        p = tf.add_paragraph()
        p.text = line
        p.level = 1
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT

    add_footer(slide)


def add_use_case_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "4) SCOR Use Case and Operating Model")

    # Three icon-style pillars
    pillars = [
        ("Modeling Objective", "Scenario Generator supports Solvency II capital modelling workflows."),
        ("Current Mode", "SCOR team runs SG software internally for calibration and output generation."),
        ("Target Operating Need", "Request to receive quarterly calibration output in CSV format."),
    ]
    for i, (title, text) in enumerate(pillars):
        x = 0.8 + i * 4.2
        icon = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.HEXAGON, Inches(x + 1.35), Inches(1.45), Inches(0.9), Inches(0.9))
        icon.fill.solid()
        icon.fill.fore_color.rgb = MOODYS_ACCENT
        icon.line.color.rgb = MOODYS_BLUE
        ip = icon.text_frame.paragraphs[0]
        ip.text = str(i + 1)
        ip.font.size = Pt(18)
        ip.font.bold = True
        ip.font.color.rgb = WHITE
        ip.alignment = PP_ALIGN.CENTER

        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.5), Inches(3.6), Inches(2.7)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0xF5, 0xF8, 0xFD)
        card.line.color.rgb = MOODYS_ACCENT

        tf = card.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = MOODYS_BLUE
        p1 = tf.add_paragraph()
        p1.text = text
        p1.font.size = Pt(12)
        p1.font.color.rgb = DARK_TEXT

    add_footer(slide)


def add_recent_topics_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "5) Recent Topics with Client (Discussion Tracker)")

    rows, cols = 5, 4
    table = slide.shapes.add_table(rows, cols, Inches(0.55), Inches(1.2), Inches(12.2), Inches(4.7)).table
    headers = ["Topic", "Client Ask", "Why It Matters", "Recommended Follow-up"]
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = MOODYS_BLUE
        para = cell.text_frame.paragraphs[0]
        para.font.bold = True
        para.font.size = Pt(12)
        para.font.color.rgb = WHITE

    items = [
        [
            "Renewal structure",
            "Proceed with single-year renewal over multi-year",
            "Indicates procurement flexibility preference",
            "Position value story for optional multi-year without forcing term commitment",
        ],
        [
            "Output delivery model",
            "Receive quarterly calibration outputs as CSV",
            "Could reduce SCOR operational effort and execution burden",
            "Assess feasibility, scope, and commercial/operational implications",
        ],
        [
            "Corporate bond parameters",
            "Clarify NumberOfBonds and Coupon settings in GenericBondPortfolios",
            "Direct impact on parameter governance and model output quality",
            "Provide technical guidance and suitable parameter boundaries",
        ],
        [
            "Dummy-value scenario",
            "Understand implications of NumberOfBonds=1, Coupon=0",
            "Potential risk of unrealistic return behavior and model distortion",
            "Run controlled test case and document impact before adoption",
        ],
    ]

    for r, row in enumerate(items, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xF8, 0xFA, 0xFD) if r % 2 == 1 else RGBColor(0xEE, 0xF3, 0xFA)
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(11)
            para.font.color.rgb = DARK_TEXT

    add_footer(slide)


def add_actions_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "6) Suggested Next Actions")

    actions = [
        "Commercial: Confirm 1-year renewal path and maintain relationship-level pricing rationale.",
        "Contract: Ensure records consistently reflect SCOR legal name/address (8 Bishopsgate).",
        "Product Ops: Define a feasibility note for quarterly CSV calibration output delivery.",
        "Quant Support: Respond formally on NumberOfBonds/Coupon settings and dummy-value impacts.",
        "Governance: Align follow-up owners across Sales, Product Specialist, and Contract teams.",
    ]

    y = 1.4
    for idx, action in enumerate(actions, start=1):
        badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.9), Inches(y), Inches(0.45), Inches(0.45))
        badge.fill.solid()
        badge.fill.fore_color.rgb = MOODYS_BLUE
        badge.line.color.rgb = MOODYS_BLUE
        bp = badge.text_frame.paragraphs[0]
        bp.text = str(idx)
        bp.font.size = Pt(12)
        bp.font.bold = True
        bp.font.color.rgb = WHITE
        bp.alignment = PP_ALIGN.CENTER

        box = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.45), Inches(y - 0.05), Inches(10.9), Inches(0.55)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xF5, 0xF8, 0xFD)
        box.line.color.rgb = MOODYS_ACCENT
        p = box.text_frame.paragraphs[0]
        p.text = action
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT

        y += 0.95

    add_footer(slide, "Source: SCOR renewal notices, amendment, and client/sales email exchange")


def build_presentation(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)
    add_executive_summary_slide(prs)
    add_contract_table_slide(prs)
    add_graph_slide(prs)
    add_use_case_slide(prs)
    add_recent_topics_slide(prs)
    add_actions_slide(prs)

    prs.save(output_path)


if __name__ == "__main__":
    build_presentation("SCOR_Moodys_Contractual_Situation_Presentation.pptx")
