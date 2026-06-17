import argparse
import os
import tempfile
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Pt


def get_layout(prs: Presentation, name: str):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    raise ValueError(f"Layout '{name}' not found in template.")


def has_layout(prs: Presentation, name: str) -> bool:
    return any(layout.name == name for layout in prs.slide_layouts)


def normalize_template_for_python_pptx(template_path: Path) -> tuple[Path, bool]:
    """
    python-pptx cannot open .potx directly (template content-type).
    Convert to a temporary .pptx package when needed.
    Returns: (loadable_path, should_cleanup)
    """
    if template_path.suffix.lower() != ".potx":
        return template_path, False

    fd, temp_path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    with zipfile.ZipFile(template_path, "r") as zin, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                )
            zout.writestr(item, data)

    return Path(temp_path), True


def set_title(slide, text: str):
    if slide.shapes.title is not None:
        slide.shapes.title.text = text


def set_placeholder_text(slide, idx: int, text: str):
    try:
        slide.placeholders[idx].text = text
    except KeyError:
        return


def style_table(table):
    # Moody's-inspired table styling with readable typography.
    header_fill = RGBColor(0x00, 0x35, 0x7A)
    header_text = RGBColor(0xFF, 0xFF, 0xFF)
    band_fill = RGBColor(0xE9, 0xF0, 0xFA)
    body_text = RGBColor(0x21, 0x2B, 0x36)

    for c in range(len(table.columns)):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = header_text

    for r in range(1, len(table.rows)):
        for c in range(len(table.columns)):
            cell = table.cell(r, c)
            if r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = band_fill
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10)
            p.font.color.rgb = body_text


def create_round_icon(path: Path, label: str, fill_rgb=(0, 53, 122)):
    img = Image.new("RGBA", (280, 280), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    draw.ellipse((10, 10, 270, 270), fill=(*fill_rgb, 255))
    text_color = (255, 255, 255, 255)
    # Use default PIL font for portability.
    bbox = draw.textbbox((0, 0), label)
    tw = bbox[2] - bbox[0]
    th = bbox[3] - bbox[1]
    draw.text(((280 - tw) / 2, (280 - th) / 2), label, fill=text_color)
    img.save(path)


def add_cover(prs: Presentation):
    slide = prs.slides.add_slide(get_layout(prs, "Cover 1"))
    set_title(slide, "SCOR Managing Agency")
    set_placeholder_text(
        slide,
        11,
        "Contractual position, use case and recent client topics\n"
        "Prepared using Moody's Scenario Generator template style",
    )
    set_placeholder_text(slide, 99, "Client: SCOR Managing Agency Limited")
    set_placeholder_text(slide, 100, "Date: June 2026")
    set_placeholder_text(slide, 101, "Prepared by: Moody's Analytics")


def add_exec_summary(prs: Presentation):
    slide = prs.slides.add_slide(get_layout(prs, "Executive Summary/Key Takeaways 3"))
    set_title(slide, "Executive summary")
    set_placeholder_text(
        slide,
        15,
        "• SCOR remains on annual renewal under Moody's agreement framework.\n"
        "• Active SG usage in SII capital modeling with tailored calibration needs.\n"
        "• Current focus: simplify operations while preserving model governance.",
    )
    set_placeholder_text(slide, 13, "Sources: 2023-2025 renewals, amendment, latest client correspondence")

    key_takeaways = {
        49: "Contract\nAnnual renewal\n(1-year preference)",
        50: "Fee trend\nGBP 53,950 in 2025\n(+3% YoY)",
        51: "Legal update\nSCOR Managing Agency\n8 Bishopsgate",
        52: "Use case\nSG supports SII\ncapital analytics",
        53: "Client ask\nQuarterly CSV\ncalibration output",
        54: "Technical topic\nNumberOfBonds/Coupon\nand dummy values",
    }
    for idx, text in key_takeaways.items():
        set_placeholder_text(slide, idx, text)


def add_contract_table(prs: Presentation):
    layout_name = "Title and Table" if has_layout(prs, "Title and Table") else "1 Column"
    slide = prs.slides.add_slide(get_layout(prs, layout_name))
    set_title(slide, "Contractual position summary")

    if layout_name == "Title and Table":
        ph = slide.placeholders[11]
        table = ph.insert_table(rows=7, cols=5).table
    else:
        container = slide.placeholders[2]
        table = slide.shapes.add_table(
            rows=7,
            cols=5,
            left=container.left,
            top=container.top,
            width=container.width,
            height=container.height,
        ).table

    headers = ["Document", "Date", "Agreement Ref", "Key point", "Fee impact"]
    for c, h in enumerate(headers):
        table.cell(0, c).text = h

    rows = [
        ["Renewal 2023", "23 Dec 2023", "00073841.7", "Annual renewal", "GBP 50,364"],
        ["Renewal 2024", "23 Dec 2024", "00073841.8", "Annual renewal", "GBP 52,379"],
        ["Renewal 2025", "23 Dec 2025", "00073841.9", "Annual renewal", "GBP 53,950"],
        ["Amendment 1", "23 Dec 2025", "00073841.9", "Legal name/address update", "No fee change"],
        ["Current term", "23 Dec 2025 - 22 Dec 2026", "ERS/SG", "Active in SII modeling", "GBP 53,950"],
        ["Commercial note", "2025 discussion", "Sales guidance", "3% uplift vs 6-10% standard", "Informational"],
    ]
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    style_table(table)


def add_fee_trend(prs: Presentation):
    slide = prs.slides.add_slide(get_layout(prs, "1 Column"))
    set_title(slide, "Commercial trend: annual renewal fees")
    set_placeholder_text(
        slide,
        17,
        "Fee progression reflects controlled uplift with 2025 increase at 3.0% year-over-year.",
    )
    set_placeholder_text(slide, 18, "Market context from sales notes: typical range ~6-10%.")

    chart_data = CategoryChartData()
    chart_data.categories = ["2023", "2024", "2025"]
    chart_data.add_series("Annual Fee (GBP)", (50364, 52379, 53950))

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        slide.placeholders[2].left,
        slide.placeholders[2].top,
        slide.placeholders[2].width,
        slide.placeholders[2].height,
        chart_data,
    ).chart
    chart.has_legend = False
    chart.series[0].format.line.color.rgb = RGBColor(0x00, 0x35, 0x7A)


def add_use_case(prs: Presentation):
    slide = prs.slides.add_slide(get_layout(prs, "2 Column, Equal - Subhead/Icons"))
    set_title(slide, "Use case and operating model")
    set_placeholder_text(slide, 41, "SG is actively used for SII capital analytics under the ERS/SG subscription.")
    set_placeholder_text(slide, 45, "Current operating model")
    set_placeholder_text(slide, 46, "Requested evolution")
    set_placeholder_text(
        slide,
        43,
        "• SCOR runs SG calibrations internally\n"
        "• Used by actuarial and capital modeling teams\n"
        "• Supports underwriting and regulatory workflows",
    )
    set_placeholder_text(
        slide,
        44,
        "• Receive quarterly calibration output in CSV\n"
        "• Clarify NumberOfBonds and Coupon guidance\n"
        "• Assess impact of dummy values (1 / 0)",
    )

    # Use template icon placeholders for cleaner Moody's visual alignment.
    csv_icon = Path(tempfile.mkstemp(suffix="_csv_icon.png")[1])
    mdl_icon = Path(tempfile.mkstemp(suffix="_mdl_icon.png")[1])
    create_round_icon(csv_icon, "NOW")
    create_round_icon(mdl_icon, "NEXT")
    try:
        slide.placeholders[37].insert_picture(str(csv_icon))
        slide.placeholders[38].insert_picture(str(mdl_icon))
    finally:
        if csv_icon.exists():
            csv_icon.unlink()
        if mdl_icon.exists():
            mdl_icon.unlink()


def add_recent_topics(prs: Presentation):
    layout_name = "Title and Table" if has_layout(prs, "Title and Table") else "1 Column"
    slide = prs.slides.add_slide(get_layout(prs, layout_name))
    set_title(slide, "Recent topics with SCOR")

    if layout_name == "Title and Table":
        ph = slide.placeholders[11]
        table = ph.insert_table(rows=5, cols=4).table
    else:
        container = slide.placeholders[2]
        table = slide.shapes.add_table(
            rows=5,
            cols=4,
            left=container.left,
            top=container.top,
            width=container.width,
            height=container.height,
        ).table

    headers = ["Topic", "Client ask", "Why it matters", "Recommended follow-up"]
    for c, h in enumerate(headers):
        table.cell(0, c).text = h

    rows = [
        ["Renewal term", "Single-year renewal", "Keeps flexibility", "Confirm term and pricing path"],
        ["CSV output", "Quarterly CSV calibration output", "Lower ops effort", "Assess delivery model"],
        ["Bond parameters", "NumberOfBonds and Coupon guidance", "Affects return realism", "Issue parameter guidance note"],
        ["Dummy values", "Effect of 1 / 0 settings", "Potential model distortion", "Run impact test and respond"],
    ]
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    style_table(table)


def add_next_steps(prs: Presentation):
    slide = prs.slides.add_slide(get_layout(prs, "5 Column - Icons"))
    set_title(slide, "Recommended next steps")
    set_placeholder_text(slide, 45, "Action plan aligned to commercial, contract, and technical follow-up")
    set_placeholder_text(slide, 46, "Source: latest SCOR correspondence and renewal documents")

    labels = ["1", "2", "3", "4", "5"]
    steps = [
        "Confirm\n1-year renewal\nexecution",
        "Validate legal\nname/address\nconsistency",
        "Assess quarterly\nCSV delivery\nmodel",
        "Issue guidance on\nbond parameters\nand dummy values",
        "Align Sales,\nProduct and\nContracts owners",
    ]

    icon_ids = [28, 36, 37, 38, 39]
    text_ids = [40, 47, 48, 49, 50]

    icon_paths = []
    for lbl in labels:
        icon_path = Path(tempfile.mkstemp(suffix=f"_{lbl}.png")[1])
        create_round_icon(icon_path, lbl)
        icon_paths.append(icon_path)

    try:
        for idx, icon_path in zip(icon_ids, icon_paths):
            slide.placeholders[idx].insert_picture(str(icon_path))
        for idx, text in zip(text_ids, steps):
            set_placeholder_text(slide, idx, text)
    finally:
        for icon_path in icon_paths:
            if icon_path.exists():
                icon_path.unlink()


def add_back_cover(prs: Presentation):
    slide = prs.slides.add_slide(get_layout(prs, "Back Cover 1"))
    set_title(slide, "Thank you")
    set_placeholder_text(slide, 15, "Questions and discussion")


def build(template_path: Path, output_path: Path, clean_only: bool = False):
    normalized_template, cleanup = normalize_template_for_python_pptx(template_path)
    prs = Presentation(str(normalized_template))
    initial_slide_count = len(prs.slides)

    try:
        add_cover(prs)
        add_exec_summary(prs)
        add_contract_table(prs)
        add_fee_trend(prs)
        add_use_case(prs)
        add_recent_topics(prs)
        add_next_steps(prs)
        add_back_cover(prs)

        if clean_only and initial_slide_count > 0:
            for i in range(initial_slide_count - 1, -1, -1):
                prs.slides._sldIdLst.remove(prs.slides._sldIdLst[i])

        prs.save(str(output_path))
    finally:
        if cleanup and normalized_template.exists():
            normalized_template.unlink()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--template",
        default="2025_Scenario_Generator_FRR (1).pptx",
        help="Path to Moody's template .pptx",
    )
    parser.add_argument(
        "--output",
        default="SCOR_Contractual_Summary_Moodys_2025_Template.pptx",
        help="Output presentation file",
    )
    parser.add_argument(
        "--clean-only",
        action="store_true",
        help="Remove template sample slides and keep only generated SCOR slides",
    )
    args = parser.parse_args()
    build(Path(args.template), Path(args.output), clean_only=args.clean_only)
