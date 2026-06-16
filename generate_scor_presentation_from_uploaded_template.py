import argparse
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


DARK_BLUE = RGBColor(0x00, 0x35, 0x7A)
LIGHT_BLUE = RGBColor(0xE9, 0xF0, 0xFA)
DARK_TEXT = RGBColor(0x21, 0x2B, 0x36)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def latest_uploaded_pptx() -> Path:
    excluded_outputs = {
        "SCOR_Contractual_Summary_Using_Uploaded_Template.pptx",
        "SCOR_Moodys_Contractual_Situation_Presentation.pptx",
    }
    candidates = sorted(
        [p for p in Path(".").glob("*.pptx") if p.name not in excluded_outputs],
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        # Fallback if only generated files exist.
        candidates = sorted(Path(".").glob("*.pptx"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not candidates:
        raise FileNotFoundError("No .pptx file found in /workspace.")
    return candidates[0]


def title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "SCOR Managing Agency"
    subtitle = slide.placeholders[1]
    subtitle.text = (
        "Contractual situation, use case, and recent discussion topics\n"
        "Prepared from SCOR renewal, amendment, and correspondence documents"
    )


def exec_summary_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Executive summary"
    tf = slide.placeholders[1].text_frame
    tf.clear()

    items = [
        "SCOR continues annual renewals under Moody's Agreement framework (originating in 2016).",
        "2025 renewal fee is GBP 53,950, a 3.0% increase vs 2024 and below referenced standard uplift range.",
        "Amendment confirms legal identity/address update to SCOR Managing Agency Limited, 8 Bishopsgate.",
        "Client preference is a one-year renewal instead of a multi-year term.",
        "Recent topics focus on CSV output delivery and corporate bond calibration parameter governance.",
    ]
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0


def contract_table_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Contractual position (summary table)"

    table = slide.shapes.add_table(6, 5, Inches(0.4), Inches(1.2), Inches(12.5), Inches(4.8)).table
    headers = ["Document", "Date", "Agreement", "Key contractual point", "Fee impact"]
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.size = Pt(11)

    rows = [
        ["Base Order Form", "23 Dec 2016", "00073841.0", "Governing baseline for subscribed SG services", "Baseline"],
        ["Renewal Notice", "23 Dec 2023", "00073841.7", "Annual renewal under same framework", "GBP 50,364"],
        ["Renewal Notice", "23 Dec 2024", "00073841.8", "Annual renewal continuation", "GBP 52,379"],
        ["Renewal Notice", "23 Dec 2025", "00073841.9", "Annual renewal continuation", "GBP 53,950"],
        ["Amendment 1", "23 Dec 2025", "00073841.9", "Entity and licensed site address updated, sanctions clause refreshed", "No explicit fee delta"],
    ]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row):
            cell = table.cell(r, c)
            cell.text = v
            if r % 2:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_BLUE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10.5)


def trend_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Commercial trend (annual renewal fees)"

    data = CategoryChartData()
    data.categories = ["2023", "2024", "2025"]
    data.add_series("Annual Fee (GBP)", (50364, 52379, 53950))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS, Inches(0.7), Inches(1.3), Inches(7.2), Inches(4.8), data
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.series[0].format.line.color.rgb = DARK_BLUE

    box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(1.7), Inches(4.4), Inches(3.9)
    )
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_BLUE
    box.line.color.rgb = DARK_BLUE
    tf = box.text_frame
    tf.clear()
    p0 = tf.paragraphs[0]
    p0.text = "Highlights"
    p0.font.bold = True
    p0.font.size = Pt(15)
    p0.font.color.rgb = DARK_BLUE
    for text in [
        "2024 vs 2023: +GBP 2,015 (+4.0%)",
        "2025 vs 2024: +GBP 1,571 (+3.0%)",
        "Sales note: market uplift often 6-10%",
        "Current pricing increase remained limited",
    ]:
        p = tf.add_paragraph()
        p.text = text
        p.level = 1
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_TEXT


def use_case_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Use case and operating model"

    pillars = [
        ("1", "Business objective", "Scenario Generator used for Solvency II capital modelling."),
        ("2", "Current mode", "SCOR runs SG software directly for calibrations and outputs."),
        ("3", "Requested evolution", "Receive calibration outputs in CSV quarterly instead of self-running."),
    ]
    left = 0.6
    for number, title, body in pillars:
        badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(left + 1.5), Inches(1.45), Inches(0.7), Inches(0.7))
        badge.fill.solid()
        badge.fill.fore_color.rgb = DARK_BLUE
        bp = badge.text_frame.paragraphs[0]
        bp.text = number
        bp.font.bold = True
        bp.font.color.rgb = WHITE
        bp.alignment = PP_ALIGN.CENTER

        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(2.3), Inches(3.9), Inches(2.8)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BLUE
        card.line.color.rgb = DARK_BLUE
        tf = card.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.bold = True
        p0.font.size = Pt(14)
        p0.font.color.rgb = DARK_BLUE
        p1 = tf.add_paragraph()
        p1.text = body
        p1.font.size = Pt(12)

        left += 4.25


def recent_topics_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Recent topics with SCOR"
    table = slide.shapes.add_table(5, 4, Inches(0.4), Inches(1.2), Inches(12.5), Inches(4.9)).table
    headers = ["Topic", "Client question / ask", "Implication", "Proposed follow-up"]
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.size = Pt(11)

    rows = [
        [
            "Renewal term",
            "Proceed with single-year renewal",
            "Signals preference for short commitment cycle",
            "Frame optional multi-year as value/operational simplification lever",
        ],
        [
            "Output format",
            "Quarterly calibration outputs in CSV",
            "Could reduce client-side operational load",
            "Assess delivery model, scope, and commercial treatment",
        ],
        [
            "Corporate bond parameters",
            "Clarify NumberOfBonds and Coupon settings",
            "Can affect return calibration behavior",
            "Provide parameter guidance note with recommended ranges",
        ],
        [
            "Dummy values",
            "Impact of NumberOfBonds=1 and Coupon=0",
            "Risk of unrealistic output and model distortion",
            "Run controlled impact test and share evidence-based response",
        ],
    ]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row):
            cell = table.cell(r, c)
            cell.text = v
            if r % 2:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_BLUE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10.5)


def next_steps_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Recommended next steps"
    tf = slide.placeholders[1].text_frame
    tf.clear()
    steps = [
        "Confirm final 1-year renewal execution path and commercial narrative.",
        "Complete contract housekeeping check for legal name and licensed address consistency.",
        "Define feasibility and ownership model for quarterly CSV output delivery.",
        "Issue formal quant response on corporate bond parameters and dummy value usage risks.",
        "Align Sales, Product Specialist, and Contracts owners for follow-through actions.",
    ]
    for i, step in enumerate(steps):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = step
        p.level = 0


def build(output_path: Path, template_path: Path | None = None) -> None:
    template_path = template_path or latest_uploaded_pptx()
    prs = Presentation(str(template_path))

    title_slide(prs)
    exec_summary_slide(prs)
    contract_table_slide(prs)
    trend_slide(prs)
    use_case_slide(prs)
    recent_topics_slide(prs)
    next_steps_slide(prs)

    prs.save(str(output_path))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=str, default=None, help="Path to uploaded .pptx template")
    parser.add_argument(
        "--output",
        type=str,
        default="SCOR_Contractual_Summary_Using_Uploaded_Template.pptx",
        help="Output presentation path",
    )
    args = parser.parse_args()
    template = Path(args.template) if args.template else None
    build(Path(args.output), template)
